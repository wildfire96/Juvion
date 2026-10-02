"""Regression checks for data persistence and editing without touching real books."""
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PyQt5.QtGui import QPixmap, QTextCursor
from PyQt5.QtWidgets import QApplication, QMessageBox, QTextEdit, QWidget

from compendium.compendium_manager import CompendiumManager
from project_window.focus_mode import FocusMode
from project_window.literary_revision_panel import LiteraryRevisionPanel
from project_window.map_panel import MapPanel
from project_window.project_model import ProjectModel
from project_window.scene_editor import SceneEditor
from project_window.templates_data import TEMPLATES, seed_template_universe
from runtime_paths import initialize_runtime
from exporter.export_dialog import ExportDialog
from exporter.export_docx import export_project_to_docx
from docx import Document
from project_window.tree_manager import save_structure
from ebooklib import epub
from exporter.export_pdf import PDFWorker
import pymupdf


class RegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.previous_cwd = Path.cwd()
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        os.chdir(self.root)
        (self.root / "Projects").mkdir()

    def tearDown(self):
        os.chdir(self.previous_cwd)
        self.directory.cleanup()

    def controller(self, model=None):
        editor = QTextEdit()
        return SimpleNamespace(model=model, scene_editor=SimpleNamespace(editor=editor), statusBar=lambda: SimpleNamespace(showMessage=lambda *_: None))

    def test_map_settings_survive_model_reload(self):
        model = ProjectModel("Mapa")
        model.settings["map"] = {"image_path": "assets/map.png", "pins": [{"entry_name": "Lugar", "x": 25, "y": 30}]}
        model.save_settings()
        self.assertEqual(ProjectModel("Mapa").settings["map"], model.settings["map"])

    def test_all_templates_create_their_universe_entries(self):
        for name, template in TEMPLATES.items():
            manager = CompendiumManager(name)
            seed_template_universe(manager, template)
            actual = manager.get_entry_index()
            self.assertEqual({e["name"] for e in actual}, {e["name"] for e in template.get("universe", [])})
            self.assertTrue(all(e["category"] not in {"Lugares", "Regras", "Itens"} for e in actual))

    def test_replacement_preserves_other_formatting_and_undo(self):
        model = SimpleNamespace(unsaved_changes=False)
        controller = self.controller(model)
        editor = controller.scene_editor.editor
        editor.setHtml("<p>Antes <b>forte</b>.</p><p>Trocar isto.</p>")
        panel = LiteraryRevisionPanel(controller)
        panel.text_source.setPlainText("Trocar isto.")
        panel.text_proposed.setPlainText("Texto revisado.")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.Yes):
            panel.apply_replacement_with_confirmation()
        self.assertIn("Texto revisado.", editor.toPlainText())
        bold = editor.document().find("forte")
        self.assertGreater(bold.charFormat().fontWeight(), 50)
        editor.undo()
        self.assertIn("Trocar isto.", editor.toPlainText())
        self.assertNotIn("Texto revisado.", editor.toPlainText())

    def test_focus_retains_rich_text(self):
        window = FocusMode(str(self.root), scene_html="<p>Um <b>nome</b>.</p>")
        saved = []
        window.on_close = saved.append
        window.close()
        self.assertEqual(len(saved), 1)
        editor = QTextEdit()
        editor.setHtml(saved[0])
        self.assertGreater(editor.document().find("nome").charFormat().fontWeight(), 50)

    def test_map_replacement_imports_image_and_clears_pins(self):
        model = ProjectModel("Mapa")
        controller = self.controller(model)
        panel = MapPanel(controller)
        image = self.root / "external.png"
        pixmap = QPixmap(40, 40)
        pixmap.fill()
        pixmap.save(str(image))
        panel.pins = [SimpleNamespace(x=lambda: 10, y=lambda: 20, entry_name="Antigo")]
        with patch("project_window.map_panel.QFileDialog.getOpenFileName", return_value=(str(image), "")), patch.object(QMessageBox, "question", return_value=QMessageBox.Yes):
            panel.select_map_image()
        self.assertEqual(model.settings["map"]["pins"], [])
        self.assertFalse(os.path.isabs(model.settings["map"]["image_path"]))
        self.assertNotIn("\\", model.settings["map"]["image_path"])
        imported = Path(panel.project_dir) / model.settings["map"]["image_path"]
        self.assertTrue(imported.exists())
        image.unlink()
        reopened = MapPanel(self.controller(ProjectModel("Mapa")))
        self.assertTrue(reopened.add_pin_btn.isEnabled())

    def test_runtime_uses_writable_directory_and_keeps_legacy_data(self):
        bundle = self.root / "bundle"
        bundle.mkdir()
        (bundle / "assets").mkdir()
        (bundle / "assets" / "resource.txt").write_text("bundled", encoding="utf-8")
        (bundle / "version.json").write_text(json.dumps({"version": "test"}), encoding="utf-8")
        (bundle / "Projects").mkdir()
        (bundle / "Projects" / "manuscript.txt").write_text("original", encoding="utf-8")
        target = self.root / "user-data"
        with patch.object(sys, "frozen", True, create=True), patch.object(sys, "_MEIPASS", str(bundle), create=True), patch.object(sys, "executable", str(bundle / "Juvion.exe")), patch.dict(os.environ, {"JUVION_DATA_DIR": str(target)}):
            initialize_runtime()
            self.assertEqual(Path.cwd(), target)
            self.assertEqual((target / "Projects" / "manuscript.txt").read_text(), "original")
            self.assertTrue((bundle / "Projects" / "manuscript.txt").exists())
            initialize_runtime()
            self.assertTrue((target / "assets" / "resource.txt").exists())
            dictionary = target / "assets" / "dictionaries"
            dictionary.mkdir()
            (dictionary / "editor_settings.json").write_text("user preference", encoding="utf-8")
            bundled_dictionary = bundle / "assets" / "dictionaries"
            bundled_dictionary.mkdir()
            (bundled_dictionary / "editor_settings.json").write_text("bundle preference", encoding="utf-8")
            (bundle / "version.json").write_text(json.dumps({"version": "next"}), encoding="utf-8")
            initialize_runtime()
            self.assertEqual((dictionary / "editor_settings.json").read_text(), "user preference")

    def test_universe_boundaries_do_not_match_partial_names(self):
        # Read the actual patterns so a duplicated escape cannot silently return.
        for filename in ("project_window/scene_editor.py", "compendium/continuity.py"):
            source = (self.previous_cwd / filename).read_text(encoding="utf-8")
            match = re.search(r'r"(\(\?<!.*?\{\}.*?\))"', source)
            self.assertIsNotNone(match)
            pattern = match.group(1).format(re.escape("Connor"))
            self.assertEqual(re.findall(pattern, "Connor, OConnor Connora Connor.", re.I), ["Connor", "Connor"])

    def test_docx_range_does_not_modify_original_structure(self):
        structure = {"acts": [{"name": "A", "chapters": [{"name": "1"}, {"name": "2"}]}, {"name": "B", "chapters": [{"name": "3"}]}]}
        combo = lambda index: SimpleNamespace(currentIndex=lambda: index)
        dialog = SimpleNamespace(project_structure=structure, start_act_combo=combo(0), end_act_combo=combo(0), start_ch_combo=combo(1), end_ch_combo=combo(1))
        selected = ExportDialog._selected_docx_structure(dialog)
        self.assertEqual(selected["acts"][0]["chapters"], [{"name": "2"}])
        self.assertEqual(len(selected["acts"]), 1)
        self.assertEqual(len(structure["acts"][0]["chapters"]), 2)

    def test_docx_export_has_real_content_without_fabricated_identifiers(self):
        path = self.root / "book.docx"
        model = SimpleNamespace(settings={}, load_scene_content=lambda *_: "<p>Olá <b>mundo</b>.</p>")
        structure = {"acts": [{"name": "Ato", "chapters": [{"name": "Capítulo", "scenes": [{"name": "Cena"}]}]}]}
        export_project_to_docx(str(path), "Livro", "Autor", structure, model=model, font_size=11.5)
        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs) + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("Olá mundo.", text)
        self.assertNotIn("000-00-00000-00-0", text)
        self.assertNotIn("Câmara Brasileira do Livro", text)
        self.assertNotIn("Dedico esta obra", text)

    def test_html_markdown_text_and_epub_exports_contain_scene(self):
        structure = {"acts": [{"name": "Ato", "chapters": [{"name": "Capítulo", "scenes": [{"name": "Cena"}]}]}]}
        save_structure("Exportar", structure)
        model = SimpleNamespace(settings={"publication": {"language": "Português"}}, load_scene_content=lambda *_: "<p>O dragão acordou.</p>")
        dialog = ExportDialog(project_name="Exportar", project_model=model)
        for kind in ("html", "markdown", "text", "epub"):
            with self.subTest(kind=kind):
                path = self.root / ("book." + kind)
                getattr(dialog, "_export_to_" + kind)(str(path), "Livro", "Autor")
                if kind == "epub":
                    book = epub.read_epub(str(path))
                    self.assertIn("O dragão acordou.", book.get_item_with_href("content.xhtml").get_content().decode("utf-8"))
                    self.assertEqual(book.get_metadata("DC", "language")[0][0], "pt-BR")
                else:
                    self.assertIn("O dragão acordou.", path.read_text(encoding="utf-8"))

    def test_pdf_export_contains_readable_scene(self):
        path = self.root / "book.pdf"
        worker = PDFWorker(str(path), "<h1>Capítulo</h1><p>O dragão acordou.</p>", "Livro", "Autor", "Arial", 12, True)
        worker.run()
        self.assertTrue(path.exists())
        with pymupdf.open(path) as document:
            extracted = "".join(page.get_text() for page in document)
            self.assertIn("O dragão acordou.", re.sub(r"\s+", " ", extracted))

    def test_spellcheck_keeps_accented_portuguese_words_whole(self):
        editor = QTextEdit()
        editor.setPlainText("coração d'água pré-história")
        checked = []
        dictionary = SimpleNamespace(lookup=lambda word: checked.append(word) or True)
        scene = SimpleNamespace(editor=editor, dictionary=dictionary, apply_extra_selections=lambda: None)
        SceneEditor.check_spelling(scene)
        self.assertEqual(checked, ["coração", "d'água", "pré-história"])


if __name__ == "__main__":
    unittest.main()
